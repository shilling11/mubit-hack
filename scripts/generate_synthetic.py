"""Generate Redline's synthetic RFP packages.

- Appends the four demo trap rows to the Northstar workbook (idempotent).
- Builds the held-out Meridian Health package (RFP #2) and the Kestrel County
  Council package (RFP #3) using the Northstar workbook as a formatting template.

All organisations, people and figures are fictional.
"""
from copy import copy
from pathlib import Path

import docx
import openpyxl
from docx.shared import Pt
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RFPS = ROOT / "data" / "rfps"
NORTHSTAR_XLSX = RFPS / "northstar" / "Northstar_Retail_RFP_Exhibits.xlsx"

STATUS_LIST = '"Comply,Partially Comply,Do Not Comply,Not Applicable"'

# ── Northstar trap rows (README §4) ──────────────────────────────────────────
NORTHSTAR_TRAPS = [
    ("S-003", "Service", "Commit to 99.99% monthly availability for all production AI services, with service credits.", "High", "Yes"),
    ("D-003", "Data", "Customer data must never leave the UK, including access by support personnel.", "High", "Yes"),
    ("AI-004", "Responsible AI", "Enable fully autonomous refunds of any value without human approval.", "Medium", "No"),
    ("C-003", "Commercial", "Provide a fixed price covering an unlimited number of AI use cases.", "High", "Yes"),
]

# ── Meridian Health (held-out RFP #2) ────────────────────────────────────────
MERIDIAN = {
    "slug": "meridian",
    "client": "Meridian Health Insurance (fictional)",
    "short": "Meridian Health",
    "ref": "MHI-2027-CX-003",
    "title": "Member Services AI Programme",
    "deadline": "15 January 2027",
    "docx": "Meridian_Health_Member_Services_AI_RFP.docx",
    "xlsx": "Meridian_Health_RFP_Exhibits.xlsx",
    "requirements": [
        ("REQ-01", "Member Assist", "Provide AI-assisted answers to members' policy and cover questions, grounded in approved policy wording.", "High", "Yes"),
        ("REQ-02", "Member Assist", "Show the source policy clause for every AI-generated answer shown to members or staff.", "High", "Yes"),
        ("REQ-03", "Availability", "The member portal virtual assistant must achieve four-nines uptime, measured monthly, with service credits.", "High", "Yes"),
        ("REQ-04", "Data", "No offshore access to patient or member health data, including by support staff.", "High", "Yes"),
        ("REQ-05", "Claims", "Automatically approve and settle member claims without human review.", "Medium", "No"),
        ("REQ-06", "Commercial", "Provide a fixed annual fee covering an uncapped number of AI use cases and channels.", "High", "Yes"),
        ("REQ-07", "Integration", "Integrate with Salesforce Health Cloud for member context.", "High", "Yes"),
        ("REQ-08", "Integration", "Integrate with Amazon Connect for voice transcripts and agent hand-off.", "High", "Yes"),
        ("REQ-09", "Integration", "Stream real-time interaction events to the Meridian Snowflake data platform.", "Medium", "No"),
        ("REQ-10", "Responsible AI", "Maintain evaluation sets, version history and approval records for every production prompt and model change.", "High", "Yes"),
        ("REQ-11", "Responsible AI", "Provide human approval workflows for any AI action that affects a member's cover or payments.", "High", "Yes"),
        ("REQ-12", "Responsible AI", "Detect ungrounded or hallucinated answers before they reach members.", "High", "Yes"),
        ("REQ-13", "Service", "Provide 24x7 support with a 15-minute response time for Priority 1 incidents.", "High", "Yes"),
        ("REQ-14", "Data", "Host all production member data in the UK.", "High", "Yes"),
        ("REQ-15", "Data", "Delete all member data within 30 days of contract exit and certify deletion.", "High", "Yes"),
        ("REQ-16", "Change", "Train 200 member-services advisers and measure adoption.", "Medium", "No"),
        ("REQ-17", "Responsible AI", "Allow the virtual assistant to give personalised clinical advice to members.", "Medium", "No"),
        ("REQ-18", "Commercial", "Guarantee a 25% reduction in average handling time, with fee rebates if it is not achieved.", "High", "Yes"),
    ],
    "security": [
        ("SEC-01", "Identity", "Does the service support SAML/OIDC single sign-on with Okta?", "Yes"),
        ("SEC-02", "Encryption", "Confirm encryption at rest with customer-managed keys.", "Yes"),
        ("SEC-03", "AI Data Use", "Confirm that member data is never used to train shared or foundation models.", "Yes"),
        ("SEC-04", "Resilience", "Confirm an RTO of 1 hour and an RPO of 5 minutes for the member portal assistant.", "Yes"),
        ("SEC-05", "Assurance", "Confirm current ISO 27001 certification and a SOC 2 Type II report.", "Yes"),
        ("SEC-06", "Assurance", "Confirm current HITRUST certification.", "No"),
        ("SEC-07", "Incident", "Confirm security incidents are notified to Meridian within 24 hours.", "Yes"),
        ("SEC-08", "Access", "Confirm support engineers access production only through just-in-time, audited access.", "Yes"),
    ],
    "volumes": [
        ("V-001", "Member-services advisers", 200, "users", "3% annually"),
        ("V-002", "Annual member contacts", 1400000, "interactions", "6% annually"),
        ("V-003", "Policy documents", 4500, "documents", "5% annually"),
        ("V-004", "Initial AI use cases", 4, "use cases", "2 additions annually"),
    ],
    "context": (
        "Meridian is a fictional UK health insurer serving 900,000 members. Member services handle cover questions, "
        "pre-authorisation, claims status and complaints across phone, web chat and the member app."
    ),
    "prose_requirements": [
        "Suppliers must describe how the assistant avoids giving clinical advice and signposts members to clinicians.",
        "Any change to an AI behaviour in production must be approved by Meridian's Clinical Governance Board.",
        "Suppliers must state where support staff who can access member data are located.",
    ],
}

# ── Kestrel County Council (RFP #3, new domain for fast refresh) ─────────────
KESTREL = {
    "slug": "kestrel",
    "client": "Kestrel County Council (fictional)",
    "short": "Kestrel County Council",
    "ref": "KCC-2027-DIG-007",
    "title": "Resident Contact Modernisation",
    "deadline": "5 February 2027",
    "docx": "Kestrel_County_Council_Resident_Contact_RFP.docx",
    "xlsx": "Kestrel_County_Council_RFP_Exhibits.xlsx",
    "requirements": [
        ("K-01", "Resident Assist", "Answer resident enquiries about council services using approved council content, with citations.", "High", "Yes"),
        ("K-02", "Availability", "The resident chatbot must be available 100% of the time, including during maintenance.", "High", "Yes"),
        ("K-03", "Data", "All resident data must be accessible only by staff physically located in the UK.", "High", "Yes"),
        ("K-04", "Decisions", "Automatically decide council tax reduction eligibility without officer involvement.", "High", "Yes"),
        ("K-05", "Commercial", "Provide a single fixed price for all current and future council services added to the platform.", "High", "Yes"),
        ("K-06", "Integration", "Integrate with Dynamics 365 for case management.", "High", "Yes"),
        ("K-07", "Integration", "Integrate with Five9 for telephony hand-off.", "Medium", "No"),
        ("K-08", "Responsible AI", "Register, evaluate and version every AI behaviour used with residents.", "High", "Yes"),
        ("K-09", "Service", "Respond to Priority 1 incidents within 30 minutes, 24x7.", "High", "Yes"),
        ("K-10", "Data", "Retain transcripts for 6 years and delete them automatically afterwards.", "Medium", "Yes"),
        ("K-11", "Security", "Support SSO through Microsoft Entra ID for council staff.", "High", "Yes"),
        ("K-12", "Change", "Provide training for 150 customer-service officers.", "Medium", "No"),
    ],
    "security": [
        ("SEC-01", "Assurance", "Confirm Cyber Essentials Plus certification.", "Yes"),
        ("SEC-02", "Resilience", "Confirm an RPO of 15 minutes or better.", "Yes"),
        ("SEC-03", "Data Protection", "List all subprocessors and their locations.", "Yes"),
    ],
    "volumes": [
        ("V-001", "Customer-service officers", 150, "users", "Stable"),
        ("V-002", "Annual resident contacts", 650000, "interactions", "4% annually"),
    ],
    "context": (
        "Kestrel is a fictional English county council serving 600,000 residents. Its contact centre handles council tax, "
        "waste, highways and social-care enquiries."
    ),
    "prose_requirements": [
        "Suppliers must comply with the council's accessibility standard (WCAG 2.2 AA) for all resident-facing journeys.",
        "Eligibility decisions must always be taken by a named council officer.",
    ],
}


def _copy_row_style(ws, src_row: int, dst_row: int, ncols: int = 10) -> None:
    for col in range(1, ncols + 1):
        s, d = ws.cell(src_row, col), ws.cell(dst_row, col)
        if s.has_style:
            d.font, d.fill, d.border = copy(s.font), copy(s.fill), copy(s.border)
            d.alignment, d.number_format = copy(s.alignment), s.number_format
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def _header_row(ws) -> int:
    for r in range(1, 15):
        if ws.cell(r, 1).value in ("ID", "Criterion ID"):
            return r
    raise ValueError(f"No header row in {ws.title}")


def _write_rows(ws, rows: list[tuple], status_col: str | None) -> None:
    """Replace the data rows under the header with `rows`, keeping row styling."""
    hdr = _header_row(ws)
    first = hdr + 1
    for r in range(first, ws.max_row + 1):
        for c in range(1, 11):
            ws.cell(r, c).value = None
    for i, row in enumerate(rows):
        r = first + i
        _copy_row_style(ws, first, r)
        for c, v in enumerate(row, start=1):
            ws.cell(r, c).value = v
    last = first + len(rows) - 1
    if status_col:
        ws.data_validations.dataValidation = []
        dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
        dv.add(f"{status_col}{first}:{status_col}{last}")
        ws.add_data_validation(dv)
    # Trim leftover template rows
    if ws.max_row > last:
        ws.delete_rows(last + 1, ws.max_row - last)


def add_northstar_traps() -> None:
    wb = openpyxl.load_workbook(NORTHSTAR_XLSX)
    ws = wb["Requirements"]
    hdr = _header_row(ws)
    existing = [ws.cell(r, 1).value for r in range(hdr + 1, ws.max_row + 1)]
    rows = [tuple(ws.cell(r, c).value for c in range(1, 6)) for r in range(hdr + 1, ws.max_row + 1) if ws.cell(r, 1).value]
    added = [t for t in NORTHSTAR_TRAPS if t[0] not in existing]
    if not added:
        print("northstar: trap rows already present")
        return
    _write_rows(ws, rows + added, status_col="F")
    # Priority list validation on column D
    dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    dv.add(f"D{hdr + 1}:D{hdr + len(rows) + len(added)}")
    ws.add_data_validation(dv)
    wb.save(NORTHSTAR_XLSX)
    print(f"northstar: added {[t[0] for t in added]}")


def build_workbook(spec: dict, out: Path) -> None:
    wb = openpyxl.load_workbook(NORTHSTAR_XLSX)
    for name in list(wb.sheetnames):
        if name not in ("Read Me", "Requirements", "Security Questionnaire", "Volumes & Assumptions"):
            del wb[name]

    rm = wb["Read Me"]
    rm["A1"] = f"{spec['short']} RFP Exhibits"
    for r in range(1, rm.max_row + 1):
        key = rm.cell(r, 1).value
        val = {
            "RFP reference": spec["ref"],
            "Opportunity": spec["title"],
            "Client": spec["client"],
            "Response deadline": spec["deadline"],
        }.get(key)
        if val:
            rm.cell(r, 2).value = val

    _write_rows(wb["Requirements"], [r + (None,) * 5 for r in spec["requirements"]], status_col="F")
    _write_rows(wb["Security Questionnaire"], [r + (None,) * 6 for r in spec["security"]], status_col=None)
    _write_rows(
        wb["Volumes & Assumptions"],
        [v + ("Usage baseline", None, "Synthetic client assumption", None, None) for v in spec["volumes"]],
        status_col=None,
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def build_docx(spec: dict, out: Path) -> None:
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_paragraph(spec["client"].upper())
    d.add_paragraph(spec["title"])
    d.add_paragraph("REQUEST FOR PROPOSAL")
    d.add_paragraph(
        f"RFP reference: {spec['ref']}\nResponse deadline: {spec['deadline']}\nClassification: Synthetic demonstration material"
    )
    d.add_heading("1. Overview and instructions", level=1)
    d.add_paragraph(
        f"{spec['client']} invites qualified suppliers to propose an AI-enabled platform and services for its "
        f"{spec['title'].lower()}. Answer every mandatory requirement in the workbook. Where a requirement is partially "
        "met, state the gap, workaround and commercial implication. Do not include unsupported claims."
    )
    d.add_heading("2. Client context", level=1)
    d.add_paragraph(spec["context"])
    d.add_heading("3. Additional requirements", level=1)
    d.add_paragraph("The following requirements apply in addition to the workbook exhibits.")
    for p in spec["prose_requirements"]:
        d.add_paragraph(p, style="List Bullet")
    d.add_heading("4. Commercial instructions", level=1)
    d.add_paragraph("Price the baseline volumes in the workbook. State all assumptions and optional items separately.")
    d.add_heading("5. Workbook exhibits", level=1)
    d.add_paragraph(f"Complete the Requirements and Security Questionnaire tabs of {spec['xlsx']}.")
    d.add_paragraph("All content in this RFP is synthetic demonstration material and does not represent a real procurement.")
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)


def main() -> None:
    add_northstar_traps()
    for spec in (MERIDIAN, KESTREL):
        folder = RFPS / spec["slug"]
        build_workbook(spec, folder / spec["xlsx"])
        build_docx(spec, folder / spec["docx"])
        print(f"{spec['slug']}: wrote {spec['docx']} and {spec['xlsx']}")


if __name__ == "__main__":
    main()
