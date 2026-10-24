"""Deterministic workbook parsing: requirement IDs come straight from the sheet."""
from pathlib import Path

import openpyxl

from app.models import Requirement

# Sheet name -> column header holding the requirement text
REQUIREMENT_SHEETS = {
    "Requirements": "Requirement",
    "Security Questionnaire": "Question",
}


def _header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, 15):
        if ws.cell(r, 1).value == "ID":
            cols = {str(ws.cell(r, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value}
            return r, cols
    raise ValueError(f"Sheet {ws.title!r} has no 'ID' header row")


def read_meta(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    meta = {"sheets": wb.sheetnames}
    if "Read Me" in wb.sheetnames:
        for row in wb["Read Me"].iter_rows(values_only=True):
            if row and row[0] and row[1]:
                meta[str(row[0]).strip()] = str(row[1]).strip()
    wb.close()
    return meta


def parse_requirements(path: Path) -> list[Requirement]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[Requirement] = []
    for sheet, text_col in REQUIREMENT_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        hdr, cols = _header(ws)
        cat_col = cols.get("Category") or cols.get("Domain")
        for r in range(hdr + 1, ws.max_row + 1):
            rid = ws.cell(r, 1).value
            text = ws.cell(r, cols[text_col]).value
            if not rid or not text:
                continue
            out.append(Requirement(
                id=str(rid).strip(),
                sheet=sheet,
                category=str(ws.cell(r, cat_col).value or "") if cat_col else "",
                text=str(text).strip(),
                priority=str(ws.cell(r, cols["Priority"]).value or "") if "Priority" in cols else "",
                mandatory=str(ws.cell(r, cols["Mandatory"]).value or "") if "Mandatory" in cols else "",
                row=r,
            ))
    return out
