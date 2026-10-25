"""Read the call exports a QA firm actually receives from a bank client, into Redress calls.

- Amazon Connect (Contact Lens) post-call JSON: one file per call. Transcript[] turns carry ParticipantId
  (AGENT / CUSTOMER), Content and Begin/EndOffsetMillis; the contact id is in CustomerMetadata.ContactId.
- A spreadsheet (Excel .xlsx or CSV), one row per line: call id, speaker, text, optional start and agent.
  This is what a call-centre system's transcript export looks like, and what QA teams work in.
- A call list spreadsheet (call id, agent, …, no text column): names each call's colleague.
- Plain text, one turn per line: "Agent: …" / "Customer: …", optionally starting "[00:12]" (a .txt file or a paste).
- Redress JSON: {id, agent_name, utterances} or a list of them.
Audio files aren't read here: they go through Deepgram like a recorded call.

Each parser returns calls as {id, agent_name, utterances: [{speaker: agent|customer, text, start_s, end_s}]}.
"""
import csv
import datetime as dt
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook

AUDIO_EXT = {".wav", ".mp3", ".m4a", ".webm", ".ogg"}
_AGENT = re.compile(r"agent|advis|colleague|\brep\b|representative|internal|staff|operator|handler", re.I)
_CUSTOMER = re.compile(r"customer|caller|client|external|member|consumer", re.I)
_GREETING = re.compile(r"thistle bank|how can i help|speaking|thanks for calling", re.I)


class ImportFailed(ValueError):
    pass


def _role(label: str) -> str | None:
    if _AGENT.search(label):
        return "agent"
    if _CUSTOMER.search(label):
        return "customer"
    return None


def _call_id(raw: str) -> str:
    return re.sub(r"[^\w-]+", "-", raw.strip()).strip("-")[:64] or "call"


def _finish(utts: list[dict]) -> list[dict]:
    """Fill missing end times from the next line's start."""
    for i, u in enumerate(utts):
        if u.get("end_s") is None:
            u["end_s"] = utts[i + 1]["start_s"] if i + 1 < len(utts) else u["start_s"] + 3
    return utts


def amazon_connect(obj: dict, stem: str) -> dict:
    roles = {p.get("ParticipantId"): p.get("ParticipantRole", p.get("ParticipantId")) for p in obj.get("Participants", [])}
    utts = []
    for turn in obj["Transcript"]:
        role = _role(str(roles.get(turn.get("ParticipantId"), turn.get("ParticipantId", ""))))
        if role and str(turn.get("Content", "")).strip():
            utts.append({"speaker": role, "text": turn["Content"].strip(),
                         "start_s": turn.get("BeginOffsetMillis", 0) / 1000, "end_s": turn.get("EndOffsetMillis", 0) / 1000})
    cid = (obj.get("CustomerMetadata") or {}).get("ContactId") or stem
    return {"id": _call_id(cid), "agent_name": "", "utterances": _finish(utts)}


def text_transcript(text: str, call_id: str) -> dict:
    line = re.compile(r"^\s*(?:\[?(?P<ts>\d{1,2}:\d{2}(?::\d{2})?)\]?\s*[-–]?\s*)?(?P<who>[A-Za-z][\w .'-]{0,30}?)\s*:\s*(?P<text>\S.*)$")
    turns: list[dict] = []
    for raw in text.splitlines():
        m = line.match(raw)
        if m:
            ts = m["ts"]
            secs = sum(int(x) * 60 ** i for i, x in enumerate(reversed(ts.split(":")))) if ts else None
            turns.append({"who": m["who"].strip(), "text": m["text"].strip(), "start_s": secs})
        elif raw.strip() and turns:
            turns[-1]["text"] += " " + raw.strip()  # a wrapped line
    if not turns:
        raise ImportFailed("no 'Speaker: text' lines found")
    # Speakers by label ("Agent", "Customer"); otherwise the one who greets is the colleague
    labels = list(dict.fromkeys(t["who"] for t in turns))
    roles = {lb: _role(lb) for lb in labels}
    if None in roles.values():
        greeter = next((t["who"] for t in turns[:3] if _GREETING.search(t["text"])), labels[0])
        for lb in labels:
            roles[lb] = roles[lb] or ("agent" if lb == greeter else "customer")
    clock, utts = 0.0, []
    for t in turns:
        start = t["start_s"] if t["start_s"] is not None else clock
        clock = start + max(1.5, len(t["text"].split()) / 2.5) + 0.4  # ~150 words a minute when there's no timing
        utts.append({"speaker": roles[t["who"]], "text": t["text"], "start_s": float(start), "end_s": None})
    agent = next((lb for lb in labels if roles[lb] == "agent" and not _AGENT.search(lb)), "")
    return {"id": _call_id(call_id), "agent_name": agent, "utterances": _finish(utts)}


def _col(headers: list[str], *names: str) -> str | None:
    norm = {re.sub(r"[^a-z]", "", h.lower()): h for h in headers}
    return next((norm[n] for n in names if n in norm), None)


def _xlsx_rows(data: bytes) -> list[dict]:
    """The first sheet as rows of strings keyed by the header row (times as m:ss, so Excel time cells work)."""
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(it, [])]
    def cell(v) -> str:
        if isinstance(v, dt.time):  # Excel turns a typed "0:03" (m:ss) into 00:03 (h:mm)
            return f"{v.hour}:{v.minute:02d}" if v.second == 0 else f"{v.hour * 60 + v.minute}:{v.second:02d}"
        if isinstance(v, (dt.datetime, dt.date)):
            return v.isoformat()
        return "" if v is None else str(v)
    return [{h: cell(v) for h, v in zip(headers, row) if h} for row in it if any(v is not None for v in row)]


def table_rows(rows: list[dict]) -> tuple[str, list[dict] | dict[str, str]]:
    """('transcript', calls) for a transcript spreadsheet, or ('call list', {call id: agent}) for a call list."""
    headers = list(rows[0].keys()) if rows else []
    id_col = _col(headers, "callid", "contactid", "call", "conversationid", "interactionid", "id")
    text_col = _col(headers, "text", "content", "transcript", "utterance", "line")
    agent_col = _col(headers, "agent", "agentname", "colleague", "advisor", "adviser")
    if not id_col:
        raise ImportFailed("no call id column (call_id or contact_id)")
    if not text_col:
        if not agent_col:
            raise ImportFailed("no text column, and no agent column for a call list")
        return "call list", {_call_id(r[id_col]): r[agent_col].strip() for r in rows if r.get(id_col)}
    speaker_col = _col(headers, "speaker", "role", "participant", "participantid", "participantrole", "party")
    start_col = _col(headers, "starts", "start", "starttime", "beginoffsetmillis", "beginoffsetms", "offsetms", "time")
    if not speaker_col:
        raise ImportFailed("no speaker column")
    calls: dict[str, dict] = {}
    for r in rows:
        if not (r.get(id_col) and (r.get(text_col) or "").strip()):
            continue
        cid = _call_id(r[id_col])
        call = calls.setdefault(cid, {"id": cid, "agent_name": "", "utterances": []})
        if agent_col and r.get(agent_col):
            call["agent_name"] = r[agent_col].strip()
        start = None
        if start_col and (r.get(start_col) or "").strip():
            v = r[start_col].strip()
            start = (sum(float(x) * 60 ** i for i, x in enumerate(reversed(v.split(":")))) if ":" in v
                     else float(v) / (1000 if "mill" in start_col.lower() or "ms" in start_col.lower() else 1))
        role = _role(r.get(speaker_col, "")) or "customer"
        call["utterances"].append({"speaker": role, "text": r[text_col].strip(), "start_s": start, "end_s": None})
    for call in calls.values():
        clock = 0.0
        for u in call["utterances"]:
            if u["start_s"] is None:
                u["start_s"] = clock
            clock = u["start_s"] + max(1.5, len(u["text"].split()) / 2.5) + 0.4
        call["utterances"].sort(key=lambda u: u["start_s"])
        _finish(call["utterances"])
    return "transcript", list(calls.values())


def parse(name: str, data: bytes) -> tuple[str, list[dict] | dict[str, str]]:
    """(format, calls), or ("call list", {call id: agent}). Raises ImportFailed with a reason."""
    stem, ext = Path(name).stem, Path(name).suffix.lower()
    text = data.decode("utf-8-sig", errors="replace")
    if ext == ".json":
        try:
            obj = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ImportFailed(f"not valid JSON ({exc.msg})") from exc
        items = obj if isinstance(obj, list) else [obj]
        if all(isinstance(o, dict) and isinstance(o.get("Transcript"), list) for o in items):
            return "Amazon Connect", [amazon_connect(o, stem if len(items) == 1 else f"{stem}-{i + 1}") for i, o in enumerate(items)]
        if all(isinstance(o, dict) and isinstance(o.get("utterances"), list) for o in items):
            return "Redress JSON", [{"id": _call_id(o.get("id") or stem), "agent_name": o.get("agent_name", ""),
                                     "utterances": _finish([dict(u) for u in o["utterances"]])} for o in items]
        raise ImportFailed("JSON, but not an Amazon Connect (Contact Lens) or Redress transcript")
    if ext in (".csv", ".xlsx"):
        try:
            rows = _xlsx_rows(data) if ext == ".xlsx" else list(csv.DictReader(io.StringIO(text)))
        except Exception as exc:  # a corrupt or password-protected workbook
            raise ImportFailed(f"couldn't read the spreadsheet ({exc})") from exc
        kind, out = table_rows(rows)
        return ("call list", out) if kind == "call list" else ("Excel" if ext == ".xlsx" else "CSV", out)
    if ext in (".txt", ".text", ".md") or not ext:
        return "text", [text_transcript(text, stem)]
    raise ImportFailed(f"unsupported file type {ext}")
