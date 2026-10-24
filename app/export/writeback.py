"""Write a run's answers back into a copy of the client's workbook."""
from pathlib import Path

import openpyxl

from app.config import EXPORTS
from app.ingest import packages
from app.ingest.xlsx_parser import REQUIREMENT_SHEETS, _header
from app.pipeline.retrieve import by_id

# Sheet -> (status column, detail column, evidence column, gap column) header names
TARGETS = {
    "Requirements": ("Supplier Response", "Response Detail", "Evidence / Proposal Ref", "Gap / Roadmap"),
    "Security Questionnaire": ("Supplier Answer", None, "Evidence / Certification", "Exception / Remediation"),
}


def export_run(run: dict, drafts: list[dict]) -> Path:
    src = packages.xlsx_path(run["package"])
    wb = openpyxl.load_workbook(src)
    kb = by_id()
    by_req = {d["req_id"]: d for d in drafts if d.get("status")}
    for sheet in REQUIREMENT_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        hdr, cols = _header(ws)
        status_h, detail_h, ev_h, gap_h = TARGETS[sheet]
        for r in range(hdr + 1, ws.max_row + 1):
            d = by_req.get(str(ws.cell(r, 1).value or "").strip())
            if not d:
                continue
            ev = "; ".join(f"{kb[e].source} — {kb[e].title}" if e in kb else e for e in d.get("evidence_ids") or [])
            if sheet == "Security Questionnaire":
                ws.cell(r, cols[status_h]).value = f"{d['status']}: {d['response']}"
            else:
                ws.cell(r, cols[status_h]).value = d["status"]
                ws.cell(r, cols[detail_h]).value = d["response"]
            ws.cell(r, cols[ev_h]).value = ev
            ws.cell(r, cols[gap_h]).value = d.get("gaps") or None
    EXPORTS.mkdir(exist_ok=True)
    out = EXPORTS / f"{src.stem}_{run['id']}.xlsx"
    wb.save(out)
    return out
